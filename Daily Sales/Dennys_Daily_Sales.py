"""PrimeLedgerAI.com - Denny's Daily Sales to QuickBooks CSV converter."""

from __future__ import annotations

import csv
import re
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

try:
    from openpyxl import load_workbook
except ImportError as exc:
    raise SystemExit("Missing openpyxl. Run: py -m pip install openpyxl") from exc


CENT = Decimal("0.01")

# Exact QuickBooks locations and cash accounts supplied by Marietta Hotel.
# The store number is used first so variations such as "5164: Berkshire" still map correctly.
LOCATION_AND_CASH = {
    "5164": ("5164-Berkshire OH", "Cash on Hand - Berkshire OH"),
    "5165": ("5165-Perrysburg OH", "Cash on Hand - Perrysburg,OH"),
    "5166": ("5166-Jeffersonville", "Cash on Hand - Jeffersonville"),
    "5167": ("5167-Elizabethtown KY", "Cash on Hand - Elizabethtown-KY"),
    "5168": ("5168-Catlettsburg KY", "Cash on hand - Catlettsburg KY"),
    "5169": ("5169-Columbus OH", "Cash on Hand - Columbus OH"),
    "7401": ("7401-Youngstown", "Cash on Hand - Youngstown"),
    "9690": ("9690 - Marietta Rest", "Cash on Hand - Marietta"),
    "9697": ("9697-Findley", "Cash on Hand - Findlay"),
}

NAMED_LOCATION_AND_CASH = {
    "marietta hotel": ("Marietta Hotel", "Cash on Hand - Marietta Hotel"),
}


def money(value) -> Decimal:
    if value is None or str(value).strip() == "":
        return Decimal("0.00")
    try:
        return Decimal(str(value).replace("$", "").replace(",", "").strip()).quantize(CENT)
    except InvalidOperation as exc:
        raise ValueError(f"Invalid dollar amount: {value!r}") from exc


def read_source(path: Path):
    rows = []
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
    else:
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb.active
        values = list(ws.iter_rows(values_only=True))
        if not values:
            return []
        headers = [str(x or "").strip() for x in values[0]]
        rows = [dict(zip(headers, row)) for row in values[1:]]

    required = {"Date", "Store", "Field", "Value"}
    if not rows and path.stat().st_size:
        raise ValueError("The input file has no data rows.")
    if rows and not required.issubset(rows[0]):
        raise ValueError("Required columns are: Date, Store, Field, Value")
    return rows


def parse_date(raw):
    text = str(raw).strip()
    if text.endswith(".0"):
        text = text[:-2]
    for fmt in ("%Y%m%d", "%m/%d/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    raise ValueError(f"Unsupported date: {raw!r}")


def location_details(store: str):
    raw = str(store).strip()
    match = re.match(r"\s*(\d+)", raw)
    if match and match.group(1) in LOCATION_AND_CASH:
        return LOCATION_AND_CASH[match.group(1)]
    named = NAMED_LOCATION_AND_CASH.get(raw.casefold())
    if named:
        return named
    valid = ", ".join(location for location, _ in LOCATION_AND_CASH.values())
    raise ValueError(f"Store {raw!r} is not in the location mapping. Valid sales locations: {valid}, Marietta Hotel")


def tender_account(field: str):
    name = field.removeprefix("Tender:").strip()
    low = name.lower()
    if "door" in low or low == "online":
        return "Credit Sales:Doordash - Sales"
    if "uber" in low:
        return "Credit Sales:Uber Eats - Sales"
    if "grub" in low:
        return "Credit Sales:Grubhub - Sales"
    if low == "qr pay":
        return "Credit Sales:QR Pay - Sales"
    if low in {"master card", "visa", "discover", "debit", "american express"}:
        return "Credit Sales:Credit card - Sales to Dennys"
    if low == "gift card":
        return "Gift Card Payable"
    if low == "cash":
        return "Cash on Hand"
    return f"Credit Sales:{name}"


def add_line(lines, account, amount, side, description, location):
    amount = money(amount)
    if amount == 0:
        return
    # Negative source values automatically reverse debit/credit.
    if amount < 0:
        side = "credit" if side == "debit" else "debit"
        amount = -amount
    lines.append({
        "Account": account,
        "Debits": amount if side == "debit" else None,
        "Credits": amount if side == "credit" else None,
        "Description": description,
        "Name": "",
        "Location": location,
    })


def build_journal(field_values, location, cash_account):
    lines = []
    get = lambda key: field_values.get(key, Decimal("0.00"))

    add_line(lines, "Sales A/C", get("Net Sales"), "credit", "Net Sales", location)
    add_line(lines, "Tax Payable:Sales Tax Payable", get("Total Sales Tax"), "credit", "Total Sales Tax", location)

    discount = get("Total Discounts")
    add_line(lines, "Food Discount A/c", discount, "debit", "Total Discounts", location)
    add_line(lines, "Food Discount A/c", discount, "credit", "Total Discounts - Offset", location)

    tips_received = get("Credit Tips Received")
    tips_paid = get("Credit Tips Paid")
    add_line(lines, "Tips Payable", tips_received, "credit", "Credit Tips Received", location)
    add_line(lines, "Tips Payable", tips_paid, "debit", "Credit Tips Paid", location)

    paid_out_non = get("Cash Paid Outs - Non Adjusting")
    paid_out_sales = get("Cash Paid Outs - Sales Adjusting")
    add_line(lines, "Cash Paid Outs", paid_out_non, "debit", "Cash Paid Outs - Non Adjusting", location)
    add_line(lines, "Cash Paid Outs", paid_out_sales, "debit", "Cash Paid Outs - Sales Adjusting", location)

    for field, amount in field_values.items():
        if not field.startswith("Tender:") or amount == 0:
            continue
        tender = field.removeprefix("Tender:").strip()
        if tender.lower() == "cash":
            # Cash tender includes tips. Tips paid and paid-outs reduce actual cash retained.
            cash_net = amount - tips_paid - paid_out_non - paid_out_sales
            add_line(lines, cash_account, cash_net, "debit", "Tender: Cash (net)", location)
        else:
            add_line(lines, tender_account(field), amount, "debit", field, location)

    debits = sum((x["Debits"] or Decimal("0")) for x in lines)
    credits = sum((x["Credits"] or Decimal("0")) for x in lines)
    variance = (debits - credits).quantize(CENT, rounding=ROUND_HALF_UP)
    if variance > 0:
        add_line(lines, "Daily Sales Variance", variance, "credit", "Auto-balance: source tender variance", location)
    elif variance < 0:
        add_line(lines, "Daily Sales Variance", -variance, "debit", "Auto-balance: source tender variance", location)
    return lines, variance


def convert(input_path: str, output_path: str):
    source = read_source(Path(input_path))
    grouped = defaultdict(dict)
    for row in source:
        if not str(row.get("Store") or "").strip() or not str(row.get("Field") or "").strip():
            continue
        key = (parse_date(row["Date"]), str(row["Store"]).strip())
        grouped[key][str(row["Field"]).strip()] = money(row.get("Value"))
    if not grouped:
        raise ValueError("No usable daily-sales records were found.")

    headers = ["Journal No.", "Journal Date", "Account", "Debits", "Credits", "Description", "Name", "Location"]
    output_rows = []

    for (date, store), fields in sorted(grouped.items()):
        location, cash_account = location_details(store)
        store_match = re.match(r"\d+", location)
        store_no = store_match.group(0) if store_match else re.sub(r"\W+", "", location)[:8]
        journal_no = f"DS{date:%m%d}-{store_no}"
        lines, variance = build_journal(fields, location, cash_account)
        for line in lines:
            output_rows.append({
                "Journal No.": journal_no,
                "Journal Date": date.strftime("%m/%d/%Y"),
                "Account": line["Account"],
                "Debits": f'{line["Debits"]:.2f}' if line["Debits"] is not None else "",
                "Credits": f'{line["Credits"]:.2f}' if line["Credits"] is not None else "",
                "Description": line["Description"],
                "Name": line["Name"],
                "Location": line["Location"],
            })

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    with Path(output_path).open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        writer.writerows(output_rows)
    return len(grouped), len(output_rows)


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("PrimeLedgerAI - Denny's Daily Sales Converter")
        self.geometry("760x280")
        self.resizable(False, False)
        self.input_var = tk.StringVar()
        self.output_var = tk.StringVar()

        ttk.Label(self, text="Denny's Daily Sales → CSV Journal", font=("Segoe UI", 18, "bold")).pack(pady=(20, 18))
        form = ttk.Frame(self, padding=12)
        form.pack(fill="x")
        ttk.Label(form, text="Input CSV/Excel:").grid(row=0, column=0, sticky="w", pady=8)
        ttk.Entry(form, textvariable=self.input_var, width=68).grid(row=0, column=1, padx=8)
        ttk.Button(form, text="Browse", command=self.choose_input).grid(row=0, column=2)
        ttk.Label(form, text="Output CSV:").grid(row=1, column=0, sticky="w", pady=8)
        ttk.Entry(form, textvariable=self.output_var, width=68).grid(row=1, column=1, padx=8)
        ttk.Button(form, text="Browse", command=self.choose_output).grid(row=1, column=2)
        ttk.Button(self, text="CREATE CSV JOURNAL", command=self.run, width=30).pack(pady=20)

    def choose_input(self):
        path = filedialog.askopenfilename(filetypes=[("Daily sales files", "*.csv *.xlsx *.xlsm"), ("All files", "*.*")])
        if path:
            self.input_var.set(path)
            p = Path(path)
            self.output_var.set(str(p.with_name(p.stem + "_Journal_Output.csv")))

    def choose_output(self):
        path = filedialog.asksaveasfilename(defaultextension=".csv", filetypes=[("CSV file", "*.csv")])
        if path:
            self.output_var.set(path)

    def run(self):
        try:
            if not self.input_var.get() or not self.output_var.get():
                raise ValueError("Please select the input and output files.")
            journals, lines = convert(self.input_var.get(), self.output_var.get())
            messagebox.showinfo("Completed", f"Created {journals} journal entries and {lines} journal lines.\n\n{self.output_var.get()}")
        except Exception as exc:
            messagebox.showerror("Could not create output", str(exc))


if __name__ == "__main__":
    App().mainloop()
