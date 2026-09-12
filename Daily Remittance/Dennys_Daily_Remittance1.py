"""PrimeLedgerAI.com - Daily Remittance CSV to Excel converter.

Select a Denny's remittance CSV export. The program creates a balanced
QuickBooks-style journal-entry workbook, grouped by restaurant unit.

Dependency (install once):  py -m pip install openpyxl
"""

from __future__ import annotations

import csv
import re
from collections import OrderedDict
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    raise SystemExit(
        "The openpyxl package is required. Open Command Prompt and run:\n"
        "py -m pip install openpyxl"
    ) from exc


# Unit-to-bank mapping. Add or change an entry here when a bank account changes.
BANK_ACCOUNTS = {
    "245164": "FCB - Sunbury ac# 7053",
    "245165": "FCB - Perrysburgh ac# 7061",
    "245166": "FCB - Jeffersonville ac# 2945",
    "245167": "FCB - ETown ac# 7045",
    "245168": "FCB - Cattletsburgh ac# 7079",
    "245169": "FCB - Columbus ac# 7037",
    "247401": "Marietta Restaurant llc (1842) - 1",
    "249690": "Marietta Restaurant llc (1842) - 1",
    "249697": "FC Bank Ac#0167 - Nights and Bites",
    # The supplied example did not identify the bank account for unit 248829.
    "248829": "BANK - Unit 8829 (UPDATE MAPPING)",
}

# QuickBooks location names. These are written on every journal line.
LOCATIONS = {
    "245164": "5164-Berkshire OH",
    "245165": "5165-Perrysburg OH",
    "245166": "5166-Jeffersonville",
    "245167": "5167-Elizabethtown KY",
    "245168": "5168-Catlettsburg KY",
    "245169": "5169-Columbus OH",
    "247401": "7401-Marietta OH",
    "248829": "8829",
    "249690": "9690",
    "249697": "9697",
}

FEE_ACCOUNT = "Commissions & fees  *:Credit Card Fee"
CARD_SALES_ACCOUNT = "Credit Sales:Credit card - Sales to Dennys"
QR_SALES_ACCOUNT = "Credit Sales:QR Pay - Sales"
HEADERS = ["DATE", "JOURNAL NO.", "ACCOUNT", "DEBITS", "CREDITS", "DESCRIPTION", "NAME", "LOCATION"]

DEFAULT_COA_RULES = [
    ("Fees Credit Cards QR", FEE_ACCOUNT, "Fees Credit Cards QR", "DEBIT"),
    ("Fees - Credit Cards", FEE_ACCOUNT, "Fees - Credit Cards", "DEBIT"),
    ("Amx Fees", FEE_ACCOUNT, "Fees - Credit Cards", "DEBIT"),
    ("Gross Credit Cards QR", QR_SALES_ACCOUNT, "Gross Credit Cards QR", "CREDIT"),
    ("Gross - Credit Cards", CARD_SALES_ACCOUNT, "Gross - Credit Cards", "CREDIT"),
]


def clean_header(value: str) -> str:
    return re.sub(r"\s+", " ", (value or "").strip()).lower()


def parse_money(value: str) -> float:
    text = (value or "").strip().replace("$", "").replace(",", "")
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    return float(text)


def excel_date(value: str) -> str:
    text = (value or "").strip()
    date_part = text[:10]
    try:
        return datetime.strptime(date_part, "%Y-%m-%d").strftime("%m/%d/%Y")
    except ValueError:
        return text


def create_mapping_workbook(path: str):
    """Create the editable unit/location/bank and chart-of-accounts mapping."""
    workbook = Workbook()
    locations_sheet = workbook.active
    locations_sheet.title = "Location Mapping"
    locations_sheet.append(["UNIT", "LOCATION", "BANK ACCOUNT"])
    for unit, location in LOCATIONS.items():
        locations_sheet.append([unit, location, BANK_ACCOUNTS.get(unit, f"BANK - Unit {unit} (UPDATE MAPPING)")])

    coa_sheet = workbook.create_sheet("Chart of Accounts")
    coa_sheet.append(["SOURCE DESCRIPTION CONTAINS", "ACCOUNT", "OUTPUT DESCRIPTION", "ENTRY TYPE"])
    for rule in DEFAULT_COA_RULES:
        coa_sheet.append(rule)

    for sheet, widths in ((locations_sheet, [16, 30, 45]), (coa_sheet, [32, 48, 32, 16])):
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        for column, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    workbook.save(path)


def load_mapping(mapping_path: str | None):
    locations = dict(LOCATIONS)
    banks = dict(BANK_ACCOUNTS)
    coa_rules = list(DEFAULT_COA_RULES)
    if not mapping_path:
        return locations, banks, coa_rules

    from openpyxl import load_workbook
    workbook = load_workbook(mapping_path, data_only=True)
    if "Location Mapping" not in workbook.sheetnames or "Chart of Accounts" not in workbook.sheetnames:
        raise ValueError("Mapping workbook must contain 'Location Mapping' and 'Chart of Accounts' sheets.")

    locations, banks = {}, {}
    for row in workbook["Location Mapping"].iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        unit = str(row[0]).strip()
        locations[unit] = str(row[1] or unit).strip()
        banks[unit] = str(row[2] or f"BANK - Unit {unit} (UPDATE MAPPING)").strip()

    coa_rules = []
    for row in workbook["Chart of Accounts"].iter_rows(min_row=2, values_only=True):
        if row[0] and row[1] and row[3]:
            coa_rules.append(tuple(str(value or "").strip() for value in row[:4]))
    if not coa_rules:
        raise ValueError("No chart-of-accounts mapping rules were found.")
    return locations, banks, coa_rules


def read_remittance(csv_path: str):
    groups = OrderedDict()
    pending_unit = None

    with open(csv_path, "r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle))
    if not rows:
        raise ValueError("The selected CSV is empty.")

    header_map = {clean_header(v): i for i, v in enumerate(rows[0])}
    required = ["deposit date", "unit", "description", "invoice amount"]
    missing = [name for name in required if name not in header_map]
    if missing:
        raise ValueError("Missing required column(s): " + ", ".join(missing))

    def field(row, name):
        index = header_map[name]
        return row[index].strip() if index < len(row) else ""

    amount_index = header_map["invoice amount"]
    for row in rows[1:]:
        unit = field(row, "unit")
        description = field(row, "description")
        deposit_date = field(row, "deposit date")

        if unit:
            pending_unit = unit
            group = groups.setdefault(unit, {"date": deposit_date, "items": [], "deposit": None})
            group["items"].append((description, parse_money(field(row, "invoice amount"))))
            continue

        # Total lines in the source place the label in Invoice amount and the
        # numeric total in the following column.
        label = row[amount_index].strip() if amount_index < len(row) else ""
        if clean_header(label).startswith("total deposit") and pending_unit:
            total_cell = row[amount_index + 1] if amount_index + 1 < len(row) else ""
            groups[pending_unit]["deposit"] = parse_money(total_cell)

    if not groups:
        raise ValueError("No remittance transactions were found.")
    return groups


def build_journal(groups, locations, banks, coa_rules):
    output = []
    warnings = []

    for unit, group in groups.items():
        date_text = excel_date(group["date"])
        try:
            date_value = datetime.strptime(date_text, "%m/%d/%Y").date()
        except ValueError:
            date_value = date_text
        try:
            journal_number = "DR" + datetime.strptime(date_text, "%m/%d/%Y").strftime("%m%d")
        except ValueError:
            journal_number = "DR"
        location = locations.get(unit, unit)
        calculated_deposit = sum(amount for _, amount in group["items"])
        deposit = group["deposit"]
        if deposit is None:
            deposit = calculated_deposit
            warnings.append(f"Unit {unit}: Total Deposit missing; calculated total was used.")
        elif abs(deposit - calculated_deposit) > 0.01:
            warnings.append(
                f"Unit {unit}: stated deposit {deposit:.2f} differs from item total "
                f"{calculated_deposit:.2f}."
            )

        for description, amount in group["items"]:
            matched = False
            for contains, account, normalized, entry_type in coa_rules:
                if contains.lower() in description.lower():
                    if entry_type.upper() == "DEBIT":
                        output.append([date_value, journal_number, account, -amount, None, f"{normalized} {date_text}", None, location])
                    elif entry_type.upper() == "CREDIT":
                        output.append([date_value, journal_number, account, None, amount, f"{normalized} {date_text}", None, location])
                    else:
                        warnings.append(f"Mapping rule '{contains}' has invalid entry type: {entry_type}")
                    matched = True
                    break
            if not matched:
                warnings.append(f"Unit {unit}: unrecognized description skipped: {description}")

        bank_account = banks.get(unit, f"BANK - Unit {unit} (UPDATE MAPPING)")
        if unit not in banks:
            warnings.append(f"Unit {unit}: bank account mapping needs to be added.")
        output.append([date_value, journal_number, bank_account, deposit, None, "BANK", None, location])

    return output, warnings


def write_excel(rows, output_path: str):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Daily Remittance"
    sheet.append(HEADERS)

    for row in rows:
        sheet.append(row)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row_number in range(2, sheet.max_row + 1):
        sheet.cell(row_number, 1).number_format = "mm/dd/yyyy"
        for column in (4, 5):
            sheet.cell(row_number, column).number_format = "#,##0.00;[Red]-#,##0.00"

    widths = [14, 16, 45, 14, 14, 46, 20, 28]
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(output_path)


def convert_file(input_path: str, output_path: str, mapping_path: str | None = None):
    groups = read_remittance(input_path)
    locations, banks, coa_rules = load_mapping(mapping_path)
    rows, warnings = build_journal(groups, locations, banks, coa_rules)
    write_excel(rows, output_path)

    debit_total = sum(float(row[3] or 0) for row in rows)
    credit_total = sum(float(row[4] or 0) for row in rows)
    if abs(debit_total - credit_total) > 0.01:
        warnings.append(
            f"Journal is not balanced: debits {debit_total:.2f}, credits {credit_total:.2f}."
        )
    return len(groups), len(rows), debit_total, credit_total, warnings


class RemittanceApp:
    def __init__(self, root):
        self.root = root
        root.title("PrimeLedgerAI.com - Daily Remittance Converter")
        root.geometry("700x355")
        root.resizable(False, False)

        frame = ttk.Frame(root, padding=22)
        frame.pack(fill="both", expand=True)
        ttk.Label(frame, text="Daily Remittance Converter", font=("Segoe UI", 18, "bold")).pack()
        ttk.Label(frame, text="PrimeLedgerAI.com", font=("Segoe UI", 10)).pack(pady=(0, 20))

        self.input_var = tk.StringVar()
        input_row = ttk.Frame(frame)
        input_row.pack(fill="x", pady=6)
        ttk.Entry(input_row, textvariable=self.input_var, width=75).pack(side="left", fill="x", expand=True)
        ttk.Button(input_row, text="Select CSV", command=self.select_csv).pack(side="left", padx=(8, 0))

        self.mapping_var = tk.StringVar()
        mapping_row = ttk.Frame(frame)
        mapping_row.pack(fill="x", pady=6)
        ttk.Entry(mapping_row, textvariable=self.mapping_var, width=75).pack(side="left", fill="x", expand=True)
        ttk.Button(mapping_row, text="Select Mapping", command=self.select_mapping).pack(side="left", padx=(8, 0))

        ttk.Button(frame, text="Create Output Excel", command=self.convert).pack(pady=22, ipadx=18, ipady=7)
        self.status = ttk.Label(frame, text="Select the daily remittance CSV file.", foreground="#555555")
        self.status.pack()

    def select_csv(self):
        path = filedialog.askopenfilename(
            title="Select Daily Remittance CSV",
            filetypes=[("CSV files", "*.csv"), ("All files", "*.*")],
        )
        if path:
            self.input_var.set(path)
            self.status.config(text=f"Selected: {Path(path).name}")

    def select_mapping(self):
        path = filedialog.askopenfilename(
            title="Select Location and Chart of Accounts Mapping",
            filetypes=[("Excel workbook", "*.xlsx"), ("All files", "*.*")],
        )
        if path:
            self.mapping_var.set(path)

    def convert(self):
        input_path = self.input_var.get().strip()
        if not input_path:
            messagebox.showwarning("Select file", "Please select the input CSV file first.")
            return

        suggested = f"{Path(input_path).stem}_QuickBooks_Journal.xlsx"
        output_path = filedialog.asksaveasfilename(
            title="Save Output Excel",
            defaultextension=".xlsx",
            initialfile=suggested,
            filetypes=[("Excel workbook", "*.xlsx")],
        )
        if not output_path:
            return

        try:
            units, row_count, debits, credits, warnings = convert_file(
                input_path, output_path, self.mapping_var.get().strip() or None
            )
            summary = (
                f"Excel created successfully.\n\nUnits: {units}\nJournal rows: {row_count}"
                f"\nDebits: {debits:,.2f}\nCredits: {credits:,.2f}\n\n{output_path}"
            )
            if warnings:
                summary += "\n\nPlease review:\n" + "\n".join(warnings)
            self.status.config(text=f"Created: {Path(output_path).name}")
            messagebox.showinfo("Conversion complete", summary)
        except Exception as exc:
            messagebox.showerror("Conversion failed", str(exc))


def main():
    default_mapping = Path(__file__).with_name("Daily_Remittance_Mapping.xlsx")
    if not default_mapping.exists():
        create_mapping_workbook(str(default_mapping))
    root = tk.Tk()
    app = RemittanceApp(root)
    app.mapping_var.set(str(default_mapping))
    root.mainloop()


if __name__ == "__main__":
    main()
