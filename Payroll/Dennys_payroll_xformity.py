"""Convert a consolidated payroll CSV to a QuickBooks payroll journal Excel file.

Run:  python payroll_to_quickbooks.py
Then select the payroll CSV. The output is saved beside the input file.
Requires: pip install openpyxl
"""
import csv
import re
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from tkinter import Tk, filedialog, messagebox

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


POSITION_MAPPING = {
    "Cook": ("Payroll:Salaries & wages:Cook & Kitchen", "Cook"),
    "Full Time Cook": ("Payroll:Salaries & wages:Cook & Kitchen", "Cook"),
    "Cook Trainee": ("Payroll:Salaries & wages:Training", "Cook Trainee"),
    "Server Trainee": ("Payroll:Salaries & wages:Training", "Server Trainee"),
    "Host/Hostess Trainee": ("Payroll:Salaries & wages:Training", "Host/Hostess Trainee"),
    "Hourly RM": ("Payroll:Salaries & wages:SalaryM", "Managers Pay"),
    "Diamond Shift Coord.": ("Payroll:Salaries & wages:Hourly Managers", "Diamond Shift Coordinator"),
    "Diamond Shift Superv": ("Payroll:Salaries & wages:Hourly Managers", "Diamond Shift Supervisor"),
    "Server": ("Payroll:Salaries & wages:Server", "Server"),
    "Premium Server": ("Payroll:Salaries & wages:Server", "Premium Server"),
    "Server-Min Wage": ("Payroll:Salaries & wages:Server", "Server-Min Wage"),
    "Host/Hostess": ("Payroll:Salaries & wages:Host/Hostess", "Host/Hostess"),
    "Service Assistant": ("Payroll:Salaries & wages:Service Assistant", "Service Assistant"),
}


def money(value):
    text = (value or "").replace("$", "").replace(",", "").strip()
    return Decimal(text or "0").quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def payroll_date(path):
    # Expected filename example: "... Jul 1-10 2026.csv"
    m = re.search(r"\b(Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)\s+(\d{1,2})(?:-\d{1,2})?\s+(\d{4})\b", path.stem, re.I)
    if not m:
        raise ValueError("Payroll start date was not found in the filename. Include a date like 'Jul 1-10 2026'.")
    return datetime.strptime(f"{m.group(1)[:3]} {m.group(2)} {m.group(3)}", "%b %d %Y")


def location_name(store):
    return store.replace(": ", "-", 1).strip()


def convert(input_path):
    entry_date = payroll_date(input_path)
    journal_no = f"Pay{entry_date:%m%d}"
    grouped = defaultdict(Decimal)
    with input_path.open(newline="", encoding="utf-8-sig") as fh:
        for row in csv.DictReader(fh):
            position = row["Position"].strip()
            account, description = POSITION_MAPPING.get(position, (f"Payroll:Salaries & wages:{position}", position))
            grouped[(location_name(row["Store"]), account, description)] += money(row["Total Wages"])

    wb = Workbook()
    ws = wb.active
    ws.title = "Journal Entry"
    headers = ["Journal No.", "Journal Date", "Account", "Debits", "Credits", "Description", "Name", "Location"]
    ws.append(headers)
    green, white = "177E75", "FFFFFF"
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=green)
        cell.font = Font(bold=True, color=white)
        cell.alignment = Alignment(horizontal="center")

    by_location = defaultdict(list)
    for (location, account, description), amount in grouped.items():
        by_location[location].append((account, description, amount))
    for location in sorted(by_location):
        location_total = Decimal("0")
        for account, description, amount in sorted(by_location[location]):
            ws.append([journal_no, entry_date, account, float(amount), None, description, "", location])
            location_total += amount
        ws.append([journal_no, entry_date, "Payroll Payable", None, float(location_total), "Payroll Total", "", location])

    widths = [16, 14, 48, 14, 14, 28, 16, 28]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        row[1].number_format = "mm/dd/yyyy"
        row[3].number_format = '$#,##0.00'
        row[4].number_format = '$#,##0.00'

    out = input_path.with_name(f"Payroll_Journal_{journal_no}.xlsx")
    wb.save(out)
    return out


def main():
    root = Tk(); root.withdraw()
    selected = filedialog.askopenfilename(title="Select consolidated payroll CSV", filetypes=[("CSV files", "*.csv")])
    if not selected:
        return
    try:
        result = convert(Path(selected))
        messagebox.showinfo("Completed", f"Payroll journal created:\n{result}")
    except Exception as exc:
        messagebox.showerror("Error", str(exc))


if __name__ == "__main__":
    main()
